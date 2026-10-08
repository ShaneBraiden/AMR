"""Data sets: the values the Studio app (entry_desktop.py) reviews, edits and confirms before forecasting.

A data set is plain JSON, so the page can edit it and the app can keep it between sessions:

    {"name": "antibiotic trend (2022-2025)", "source": "upload" | "manual", "file": "antibiotic trend (2022-2025).xlsx",
     "years": [2022, 2023, 2024, 2025],
     "organisms": [{"name": "Escherichia coli", "note": null,
                    "rows": [{"medicine": "Meropenem", "original": "Meropenam", "values": [18.6, 6.82, 11.36, 31.91]}]}]}

Values are percentages (0-100). A cell that can't be used keeps what it held - null when empty, the text otherwise -
so the review can show it; problems() lists everything to fix before forecasting.

A confirmed data set is written as a workbook in the same layout as the source file (write_workbook), which the
forecasting engine reads like any other. Its values come out of read_workbook exactly as they went in, so a workbook
uploaded and confirmed unchanged gives the same forecasts as the workbook itself.
"""

import hashlib
import json
import re
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from forecast_engine import CACHE_VERSION
from parse_antibiotic_trend import (ANTIBIOTIC_NAMES, MIN_YEARS, ORGANISM_NAMES, WorkbookError, find_header, load,
                                    organism_name, year_columns)

SHEET_NAME_LIMIT = 31  # Excel's limit; an organism's name is its sheet's name
SHEET_NAME_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")
SHEET_NAMES = {name: sheet for sheet, name in ORGANISM_NAMES.items()}  # "Escherichia coli" -> "E.COLI"
PERCENT_TEXT = re.compile(r"^\s*(-?\d+(?:\.\d+)?)\s*%\s*$")
NO_DATA = "No data recorded."


def pct(value: float) -> float:
    """Round a percentage to 2 places the way read_workbook does, so the engine sees the same numbers."""
    return float(np.round(float(value), 2))


def cell_value(v):
    """A cell as the review shows it: a percentage, null when empty, or the text that was there."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return pct(v * 100)  # Excel stores 25% as 0.25
    match = PERCENT_TEXT.match(str(v))
    return pct(match.group(1)) if match else str(v).strip()


def from_workbook(path: Path) -> dict:
    """Read every value of a workbook for review, keeping cells that need fixing instead of rejecting the file.

    Raises WorkbookError only for a layout this can't read: no header row, or year columns that aren't consecutive
    or differ between sheets.
    """
    path = Path(path)
    wb = load(path)
    years, organisms = None, []
    for ws in wb.worksheets:
        header_row, header = find_header(ws)
        sheet_years, _net = year_columns(ws, header)
        if years is None:
            years = sheet_years
        elif sheet_years != years:
            raise WorkbookError(f"Sheet '{ws.title}' has years {sheet_years[0]}–{sheet_years[-1]} but the first sheet "
                                f"has {years[0]}–{years[-1]}; every sheet needs the same years.")
        rows, note = [], None
        for cells in ws.iter_rows(min_row=header_row + 1, values_only=True):
            cells = list(cells) + [None] * (1 + len(years) - len(cells))
            name, values = cells[0], cells[1: 1 + len(years)]
            name = "" if name is None else str(name).strip()
            if all(cell_value(v) is None for v in values):
                if name:
                    note = name  # a row with text but no numbers is a note (e.g. Streptococcus has no data)
                continue
            rows.append({"medicine": ANTIBIOTIC_NAMES.get(name, name), "original": name,
                         "values": [cell_value(v) for v in values]})
        organisms.append({"name": organism_name(ws.title), "note": note, "rows": rows})
    return {"name": path.stem, "source": "upload", "file": path.name, "years": years, "organisms": organisms}


def sheet_name(organism: str) -> str:
    return SHEET_NAMES.get(organism, organism)


def number(v) -> float | None:
    """A cell's percentage, or None if it isn't a number."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    if isinstance(v, str):
        match = PERCENT_TEXT.match(v) or re.match(r"^\s*(-?\d+(?:\.\d+)?)\s*$", v)
        if match:
            return float(match.group(1))
    return None


def problems(dataset: dict) -> list[dict]:
    """Everything to fix before the data set can be forecast; empty when it is ready.

    Each problem says where it is - "where" is dataset, years, organism, medicine or value - with the organism's
    index, the row's index and the year's index as they apply.
    """
    found = []

    def add(message, where="dataset", organism=None, row=None, year=None):
        found.append({"where": where, "organism": organism, "row": row, "year": year, "message": message})

    years = dataset.get("years") or []
    if not all(isinstance(y, int) and 1000 <= y <= 9999 for y in years):
        add("Years must be written in full, like 2022.", "years")
    elif len(years) < MIN_YEARS:
        add(f"Add at least {MIN_YEARS} years: the models need two years to learn from, one to predict and one "
            f"to check that prediction against.", "years")
    elif years != list(range(years[0], years[0] + len(years))):
        add("The years must follow on from each other.", "years")

    organisms = dataset.get("organisms") or []
    if not any(org.get("rows") for org in organisms):
        add("Add at least one medicine with its percentages.")

    seen_organisms = {}
    for i, org in enumerate(organisms):
        name = organism_name(org.get("name") or "")
        if not name:
            add("Name this organism.", "organism", i)
        elif SHEET_NAME_FORBIDDEN.search(name):
            add("Organism names can't contain [ ] : * ? / or \\ (Excel doesn't allow them in sheet names).",
                "organism", i)
        elif len(sheet_name(name)) > SHEET_NAME_LIMIT:
            add(f"Organism names can be at most {SHEET_NAME_LIMIT} characters (an Excel limit).", "organism", i)
        elif name.lower() in seen_organisms:
            add(f"{name} is listed twice.", "organism", i)
        else:
            seen_organisms[name.lower()] = i

        seen_medicines = set()
        for j, row in enumerate(org.get("rows") or []):
            raw = (row.get("medicine") or "").strip()
            medicine = ANTIBIOTIC_NAMES.get(raw, raw)
            if not medicine:
                add("Name this medicine.", "medicine", i, j)
            elif medicine.lower() in seen_medicines:
                add(f"{medicine} is listed twice for this organism.", "medicine", i, j)
            else:
                seen_medicines.add(medicine.lower())

            values = list(row.get("values") or [])
            for k in range(len(years)):
                v = values[k] if k < len(values) else None
                n = number(v)
                if v is None or (isinstance(v, str) and not v.strip()):
                    add("Enter a percentage.", "value", i, j, k)
                elif n is None:
                    add(f"“{v}” is not a number.", "value", i, j, k)
                elif not 0 <= n <= 100:
                    add(f"{n:g}% is outside 0–100%.", "value", i, j, k)
    return found


def canonical(dataset: dict) -> dict:
    """The data set as it will be forecast: names as the dashboard shows them, values as rounded numbers.

    Only for a data set without problems.
    """
    out = {"years": list(dataset["years"]), "organisms": []}
    for org in dataset["organisms"]:
        rows = []
        for row in org.get("rows") or []:
            raw = row["medicine"].strip()
            rows.append({"medicine": ANTIBIOTIC_NAMES.get(raw, raw), "values": [pct(number(v)) for v in row["values"]]})
        note = None if rows else ((org.get("note") or "").strip() or NO_DATA)
        out["organisms"].append({"name": organism_name(org["name"]), "note": note, "rows": rows})
    return out


def dataset_key(dataset: dict) -> str:
    """Cache name for a data set: changes whenever its values or the models change, not with its name."""
    text = json.dumps(canonical(dataset), sort_keys=True, ensure_ascii=False)
    return f"v{CACHE_VERSION}-d{hashlib.sha256(text.encode()).hexdigest()[:24]}"


def write_workbook(dataset: dict, path: Path) -> Path:
    """Write a data set without problems as a workbook in the source file's layout: a sheet per organism."""
    data = canonical(dataset)
    years = data["years"]
    wb = Workbook()
    wb.remove(wb.active)
    header = ["Antibiotic", *[f"{y} (%)" for y in years], "Net Change (%)"]
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    listed = ", ".join(map(str, years[:-1])) + f", and {years[-1]}"
    for org in data["organisms"]:
        ws = wb.create_sheet(sheet_name(org["name"]))
        ws["A1"] = f"{org['name']} - Resistance Progression ({years[0]} to {years[-1]})"
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = f"Tracking resistance percentages across all intermediate years ({listed})."
        ws["A2"].font = Font(italic=True)
        for r in (1, 2):
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(header))
        for col, text in enumerate(header, start=1):
            cell = ws.cell(row=4, column=col, value=text)
            cell.font, cell.fill = bold, fill
            cell.alignment = Alignment(horizontal="left" if col == 1 else "center")
        for r, row in enumerate(org["rows"], start=5):
            v = row["values"]
            ws.cell(row=r, column=1, value=row["medicine"])
            for col, value in enumerate([*v, v[-1] - v[0]], start=2):
                ws.cell(row=r, column=col, value=value / 100).number_format = "0.00%"
        if org["note"]:
            ws.cell(row=5, column=1, value=org["note"])
        ws.column_dimensions["A"].width = max([28, *(len(row["medicine"]) + 2 for row in org["rows"])])
        for col in range(2, len(header) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 15
        ws.freeze_panes = "B5"
    wb.save(path)
    return Path(path)
